"""唯讀驗收實際三份成果，包含 XLSX 已重算快取及跨文件需求對映。"""
import json
from math import isclose
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/revised_20260908'


def check(condition,message):
    if not condition:
        raise AssertionError(message)


def near(actual,expected,label):
    check(isinstance(actual,(int,float)) and isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-8),f'{label}: {actual} != {expected}')


def main():
    package = json.loads((OUT/'design_package.json').read_text(encoding='utf-8'))
    check(not list(OUT.glob('*.pdf')),'本輪交付不應包含 PDF')
    books = {}
    formula_count = 0
    for filename in ['設計計算表.xlsx','設備明細表.xlsx']:
        values = openpyxl.load_workbook(OUT/filename,data_only=True)
        formulas = openpyxl.load_workbook(OUT/filename,data_only=False)
        books[filename] = values
        for sheet in formulas:
            check(bool(sheet.print_area),f'{sheet.title} 沒有列印範圍')
            check(bool(sheet.print_title_rows),f'{sheet.title} 沒有重複欄頭')
            check(sheet.page_setup.orientation=='landscape' and sheet.page_setup.fitToWidth==1,f'{sheet.title} 列印寬度設定')
            for row in sheet:
                for cell in row:
                    cached = values[sheet.title][cell.coordinate]
                    check(cached.data_type!='e',f'{sheet.title}!{cell.coordinate} 公式錯誤')
                    if cell.data_type=='f':
                        formula_count += 1
                        check(isinstance(cached.value,(int,float)),f'{sheet.title}!{cell.coordinate} 缺重算數值快取')
        formulas.close()
    calc = books['設計計算表.xlsx']; equip = books['設備明細表.xlsx']
    near(calc['概覽']['C20'].value,920,'總新風 L/s')
    near(calc['概覽']['D20'].value,87.5,'總排風 L/s')
    near(calc['通風計算']['M15'].value,920,'房間新風合計')
    near(calc['通風計算']['J15'].value,87.5,'房間排風合計')
    near(calc['通風計算']['P15'].value,3312,'新風 m³/h')
    near(calc['通風計算']['I15'].value,315,'排風 m³/h')
    for row in range(5,8):
        sheet = calc['系統負荷密度']
        area,raw,sf,sizing,density,sizing_density = [sheet.cell(row,c).value for c in range(2,8)]
        near(sf,1.15,'SF')
        near(sizing,raw*1.15,'單次 SF')
        near(density,raw*1000/area,'原始密度 W/m²')
        near(sizing_density,sizing*1000/area,'含 SF 密度 W/m²')
    for row in range(5,14):
        sheet = calc['冷負荷計算']
        near(sheet.cell(row,9).value,sheet.cell(row,7).value*1000/sheet.cell(row,3).value,'房間負荷密度')
        check(sheet.cell(row,6).value=='未計算','新風盤管缺值被改寫')
    def all_text(book):
        return '\n'.join(str(c.value) for s in book for row in s for c in row if c.value is not None)
    calc_text,equip_text = all_text(calc),all_text(equip)
    with ZipFile(OUT/'設計說明.docx') as z:
        doc = ET.fromstring(z.read('word/document.xml'))
    word_text = ''.join(n.text or '' for n in doc.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    expected = {e['id'] for e in package['equipment']}
    actual = {str(row[0].value) for s in equip for row in s.iter_rows(min_row=8) if row[0].value in expected}
    check(actual==expected,'設備表需求編號不完整')
    for identifier in expected:
        check(identifier in calc_text and identifier in word_text,f'{identifier} 跨文件對映缺漏')
    for e in package['equipment']:
        s = equip[{'AC':'空調設備','TEF':'風機','FAU':'新風處理'}[e['category']]]
        row = next(r for r in s if r[0].value==e['id'])
        check(row[3 if e['category']=='AC' else 2].value=='待確定','未知數量不可變成0或1')
        for rid in e['room_ids']:
            check(rid in row[1].value,f'{e["id"]} 遺漏房間 {rid}')
    for unit in ['m²','m³','(人)','L/s','m³/h','kW','W/m²','Pa','ACH']:
        check(unit in calc_text,f'計算表缺單位 {unit}')
    for unit in ['台','kW','L/s','Pa','rpm','V/Ph/Hz','dB(A)','mm','°C DB']:
        check(unit in equip_text,f'設備表缺單位 {unit}')
    check(formula_count>0,'沒有公式可驗收')
    for book in books.values():book.close()
    print(f'通過：三份成果、{formula_count} 個公式快取、{len(expected)} 個需求跨文件對映、風量／SF／密度／單位／列印設定。')


if __name__=='__main__':
    main()
