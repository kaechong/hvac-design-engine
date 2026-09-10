"""核對四份成果與單機、跟進及實核資料。"""
import json
import os
from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=Path(os.environ.get('HVAC_OUTPUT_DIR',ROOT/'outputs/drawing_handoff_20260910'))

def main():
    env=os.environ.copy();env['HVAC_OUTPUT_DIR']=str(OUT);env['PYTHONIOENCODING']='utf-8'
    subprocess.run([sys.executable,str(ROOT/'scripts/verify_selected_outputs.py')],env=env,check=True,cwd=ROOT)
    p=json.loads((OUT/'design_package.json').read_text(encoding='utf-8'))
    md=(OUT/'畫圖工作指引.md').read_text(encoding='utf-8')
    units=p['drawing_equipment'];ids=[u['equipment_id'] for u in units]
    assert len(ids)==len(set(ids))==sum(c['quantity'] for c in p['configuration'])==17
    book=load_workbook(OUT/'設備明細表.xlsx',data_only=True)
    sheet=book['單機定位'];values='\n'.join(str(c.value) for row in sheet for c in row if c.value is not None)
    for u in units:
        assert '### '+u['equipment_id']+'　' in md
        assert u['equipment_id'] in values and u['model'] in values
        assert u['group_id'] in md and u['source_pages']
        assert u['unresolved_impact'] and u['service_access']
    book.close()
    calc=load_workbook(OUT/'設計計算表.xlsx',data_only=True)
    content='\n'.join(str(c.value) for s in calc for row in s for c in row if c.value is not None)
    for a in p['review_actions']:
        for key in ('current_adopted','steps','required_evidence','provisional_action','completion_criteria','calculation_impact','drawing_impact'):
            assert a[key],(a['id'],key)
        assert a['id'] in content and a['id'] in md
    assert len(p['fop_actual_checks'])==6
    for check in p['fop_actual_checks']:
        assert check['status']=='待核實' and check['evidence']==[]
        assert check['item'] in content
    calc.close()
    assert not list(OUT.glob('*.pdf'))
    assert next(u for u in units if u['equipment_id']=='IU-GF-02')['proposed_location'] is None
    assert all(u['proposed_location'] is None for u in units if u['role']=='室外機')
    print('通過：四份成果、17台單機、10項可執行跟進、6項獨立實核；來源未明的安裝位置保持空值。')

if __name__=='__main__':main()
