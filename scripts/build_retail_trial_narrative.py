"""從 FOP 原母本複製版式，固定章原文經版本化模板逐段核驗。"""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/fop_narrative_v1.1.json"


def full_text(element):
    return ''.join(t.text or '' for t in element.iter(qn('w:t')))


def replace_text(element, text):
    """保留段落和內容控制結構；只替換文字節點，不重建母本樣式。"""
    nodes = list(element.iter(qn('w:t')))
    if not nodes:
        target = element
        if element.tag == qn('w:tc'):
            target = element.find(qn('w:p'))
            if target is None:
                target = OxmlElement('w:p')
                element.append(target)
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        r.append(t)
        target.append(r)
        nodes = [t]
    nodes[0].text = str(text)
    for node in nodes[1:]:
        node.text = ''


def populate_table(table, headers, rows):
    samples = list(table._tbl.iterchildren(qn('w:tr')))
    header, body = deepcopy(samples[0]), deepcopy(samples[1])
    for row in samples:
        table._tbl.remove(row)
    for index, values in enumerate([headers, *rows]):
        row = deepcopy(header if index == 0 else body)
        # 母本 Google 內容控制具有唯一 ID；複製房間列不可複製識別碼。
        for node in list(row.iter()):
            if node.tag in (qn('w:bookmarkStart'), qn('w:bookmarkEnd'), qn('w:id')):
                node.getparent().remove(node)
        cells = list(row.iterchildren(qn('w:tc')))
        for cell, value in zip(cells, values):
            replace_text(cell, value)
        props = row.find(qn('w:trPr'))
        if props is None:
            props = OxmlElement('w:trPr')
            row.insert(0, props)
        for tag in ('w:trHeight', 'w:cantSplit', 'w:tblHeader'):
            for existing in list(props.findall(qn(tag))):
                props.remove(existing)
        props.append(OxmlElement('w:cantSplit'))
        if index == 0:
            props.append(OxmlElement('w:tblHeader'))
        table._tbl.append(row)


def build_document(package, reference=None):
    contract = json.loads(TEMPLATE.read_text(encoding='utf-8'))
    reference = reference or next((ROOT / 'data/local').rglob('*Design_Narrative_v1.1.docx'))
    expected = next(x['sha256'] for x in contract['source_files'] if x['name'].endswith('.docx'))
    if hashlib.sha256(reference.read_bytes()).hexdigest() != expected:
        raise ValueError('FOP 母本 SHA256 不符，須重新覆核版本。')
    doc = Document(reference)
    for paragraphs in contract['fixed_chapters'].values():
        for item in paragraphs:
            element = doc.paragraphs[item['paragraph_index']]._p
            if full_text(element) != item['text']:
                raise ValueError('固定章原文與版本化模板不一致。')
            replace_text(element, item['text'])

    inp, result = package['input'], package['result']
    rooms = {r['room_id']: r for r in inp['rooms']}
    outcomes = {r['room_id']: r for r in result['rooms']}
    equipment = package['equipment']
    configurations = package.get('configuration', [])
    treatment = package.get('air_treatment')
    if not treatment:
        raise ValueError('正式報告須提供新風盤管及再熱計算。')
    area = sum(r['room_metadata']['area_m2'] for r in rooms.values())
    cooling = sum(s['sizing_peak']['total_kw'] for s in result['systems'])
    raw = sum(s['raw_peak']['total_kw'] for s in result['systems'])
    fresh = sum(s['fresh_air_ls'] for s in result['systems'])
    exhaust = sum(s['exhaust_ls'] for s in result['systems'])
    def refs(room_id, prefix):
        if prefix=='AC-' and rooms[room_id]['exhaust_ach']>0:
            return '—'
        return '、'.join(e['id'] for e in equipment if room_id in e['room_ids'] and e['id'].startswith(prefix)) or '—'
    def models_for(room_id):
        found = [f"{e['equipment_id']}\n{e['quantity']} 台" for e in configurations
                 if room_id in e.get('room_ids', []) and e.get('role') in ('indoor', 'indoor_unit', 'dx_indoor', 'terminal', '室內機')]
        return '、'.join(found) or refs(room_id, 'AC-')
    model_list = '；'.join(dict.fromkeys(f"{e.get('brand', '')} {e['model']}" for e in configurations if e.get('model')))
    coil_raw = treatment['raw_coil']['total_kw']
    coil_sizing = treatment['sizing_coil']['total_kw']
    coil_out = treatment['coil_outlet']
    values = {
        0: '機械通風和空調系統設計說 明 及 解 釋 備 忘 錄',
        1: f"零售展廳三層空調通風工程\n版本：{package['version']}",
        3: '1.1 項目名稱：零售展廳地庫、地面層及一樓空調通風工程',
        4: '',
        5: f'1.2 工程範圍：本工程涵蓋地庫、地面層及一樓共九個房間，計算面積 {area:.2f} m²，設置直接膨脹式空調末端、新風冷卻除濕及再熱系統，以及衛生間機械排風。',
        39: '• 室外設計氣象條件 (Outdoor Design Temperature)：澳門夏季。',
        40: '• 夏季 (Summer)：34°C 乾球／28°C 濕球；室內 23°C、55% RH。',
        41: '',
        42: '• 室內空間設計參數與負荷摘要：下表列出房間原始峰值冷負荷；冷負荷密度（W/m²）＝原始峰值冷負荷（kW）× 1,000／房間面積（m²）。',
        44: '3.2.1 空調系統與直接膨脹式冷源規範 (DX Provision)',
        45: f'本工程按樓層及服務區域配置直接膨脹式空調。室內末端原始冷負荷合計 {raw:.2f} kW，設計冷量合計 {cooling:.2f} kW；安全係數 1.15 於各系統原始尖峰冷負荷施加一次。按面積 {area:.2f} m² 計算，原始及設計冷負荷密度分別為 {raw*1000/area:.2f} 及 {cooling*1000/area:.2f} W/m²。上述為末端系統需求合計，新風處理盤管冷量另列。設備候選包括 {model_list}；配置及性能核對狀態見設備明細表。',
        47: '設計新風量採 10 L/s／人，衛生間排風採 15 h⁻¹（ACH）。新排風量不施加冷負荷安全係數。',
        48: '',
        49: f'• 鮮風供應：設計新風量 {fresh:.2f} L/s（{fresh*3.6:.2f} m³/h），設備需求編號 FAU-REQ-ALL。室外新風經盤管冷卻除濕至 {coil_out["db_c"]:.2f}°C、{coil_out["rh_fraction"]*100:.0f}% RH，再熱至室內設計送風狀態 23°C、55% RH。盤管原始冷負荷 {coil_raw:.2f} kW，含安全係數 1.15 的選型冷量 {coil_sizing:.2f} kW；再熱計算需求 {treatment["reheat_kw"]:.2f} kW。末端與新風盤管設計需求分列，盤管冷量不重複計入末端冷量；兩類設計需求合計 {cooling+coil_sizing:.2f} kW。',
        50: f'• 排風：地庫及地面層衛生間合計排風量 {exhaust:.2f} L/s（{exhaust*3.6:.2f} m³/h），對應 TEF-REQ-BF／TEF-REQ-GF。排風量（m³/h）＝體積（m³）×換氣次數（h⁻¹）。補風由鄰接空間流向衛生間，維持衛生間相對負壓。',
        53: '• 設備安裝：壁掛式室內機須核對牆體承載、送回風間距及檢修空間；吊裝風機的吊架按運行重量設計並設隔振裝置。設備安裝高度須與室內及天花配置協調。',
        54: '• 軟連接：風管與風機／風管式室內機連接處採柔性連接，規格由正式設備及安裝要求確定。',
        55: '• 空調末端設備：選型須同時滿足設計工況下總冷量、顯熱及潛熱需求，並核對室內外機配對、最低調節能力、風量及噪音。候選型號須完成原廠工況性能核對後定案。',
        56: '• 風管與閘閥：風管按設計風量及阻力確定尺寸，風機按同一轉速的風量與靜壓工作點選定。穿越防火分區處設相應防火閥及檢修口，並配合消防聯鎖要求。',
        57: '• 管道與保溫：設置冷媒及凝結水管；冷媒管管徑、長度及高差須符合原廠配對限制，保溫按防結露要求配置。凝結水管設排水坡度及必要的存水彎。',
        59: '• 動力和控制材料依本項目電氣技術規範選定。',
        60: '• 設備設獨立隔離開關，電纜及保護裝置按設備電氣參數配合。輸入功率（kW）及電源（V/Ph/Hz）詳見設備明細表。',
        61: '• 溫濕度控制：室內設定 23°C、55% RH；新風機組按盤管離風狀態控制除濕，並以再熱控制送風溫度。末端依各服務區域負荷調節。',
        62: '• 火警連動：設備停機及閥門聯鎖按消防分區與控制接口配置。',
    }
    for index, text in values.items():
        replace_text(doc.paragraphs[index]._p, text)
    doc.paragraphs[0].style = doc.styles['Title']
    for holder in (doc.paragraphs[0]._p.get_or_add_pPr(), doc.styles['Title'].element.get_or_add_pPr()):
        for border in list(holder.findall(qn('w:pBdr'))):
            holder.remove(border)
    for index in [2,7,8,19,37,38,43,44,46,51,52,58,63,70,72,74]:
        doc.paragraphs[index].paragraph_format.keep_with_next = True
    # 對照母本頁面設定；不增加封面或額外裝飾。
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    spec = json.loads((ROOT/'data/reference/fop_template_spec.json').read_text(encoding='utf-8'))['docx']
    populate_table(doc.tables[0], spec['tables'][0]['rows'][0], [
        [rid, r['room_metadata']['name_zh'], r['room_metadata']['floor'], f"{r['room_metadata']['area_m2']:.2f}", f"{r['volume_m3']/r['room_metadata']['area_m2']:.2f}", r['room_metadata']['name_zh']]
        for rid, r in rooms.items()])
    summary = []
    for rid, r in rooms.items():
        m, o = r['room_metadata'], outcomes[rid]
        summary.append([f"{rid}\n{m['name_zh']}", f"{m['area_m2']:.2f}", str(r['people']), f"{o['ventilation']['exhaust_m3h']:.2f}", f"{o['ventilation']['fresh_air_ls']:.2f}", f"{o['raw_peak']['total_kw']:.2f}", f"{o['raw_peak']['total_kw']*1000/m['area_m2']:.2f}", models_for(rid)])
    populate_table(doc.tables[1], ['房間名稱','面積\n(m²)','人數\n(人)','排風\n(m³/h)','新風\n(L/s)','原始峰值\n(kW)','原始密度\n(W/m²)','本工程空調通風配置'], summary)
    populate_table(doc.tables[2], ['樓層','區域名稱','空間用途','空調需求記錄','新風需求記錄','排風需求記錄'], [
        [r['room_metadata']['floor'], rid+'\n'+r['room_metadata']['name_zh'],r['room_metadata']['name_zh'],refs(rid,'AC-'), refs(rid,'FAU-'),refs(rid,'TEF-')]
        for rid,r in rooms.items()])
    reasons=['核對設備明細與各專業配置的一致性','核對設備、管道與逃生及隔火空間','核對送回風配置與室內噪音','核對設備安裝與消防要求','核對管道防火穿越及封堵','核對設備安裝高度與天花淨空']
    populate_table(doc.tables[3], ['項次','內容','核實狀態','證據／待辦'], [
        [row[0],row[1],'待核實',reasons[i]] for i,row in enumerate(spec['tables'][3]['rows'][1:])])
    # 完成依原母本段落編號填寫後移除不適用段落。
    for index in sorted((4, 41, 48), reverse=True):
        element = doc.paragraphs[index]._p
        element.getparent().remove(element)
    for paragraph in doc.paragraphs:
        if full_text(paragraph._p).startswith('5 –'):
            paragraph.paragraph_format.page_break_before = True
    return doc


def main():
    parser = argparse.ArgumentParser()
    output_dir=Path(os.environ.get('HVAC_OUTPUT_DIR',ROOT/'outputs/selected_20260908'))
    parser.add_argument('--package', type=Path, default=output_dir/'design_package.json')
    parser.add_argument('--output', type=Path, default=output_dir/'設計說明.docx')
    args = parser.parse_args()
    package = json.loads(args.package.read_text(encoding='utf-8'))
    doc = build_document(package)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == '__main__':
    main()
