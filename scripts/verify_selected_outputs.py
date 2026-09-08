"""唯讀驗收三份 r3 成果：公式快取、來源、單機配置、固定文案及負荷邊界。"""
import json
import os
from math import isclose
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT=Path(os.environ.get('HVAC_OUTPUT_DIR',ROOT/'outputs/selected_20260908'))
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def check(condition,message):
    if not condition:raise AssertionError(message)


def near(actual,expected,label):
    check(isinstance(actual,(int,float)) and isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-8),f'{label}: {actual} != {expected}')


def all_text(book):
    return '\n'.join(str(c.value) for s in book for row in s for c in row if c.value is not None)


def main():
    package=json.loads((OUT/'design_package.json').read_text(encoding='utf-8'))
    check(not list(OUT.glob('*.pdf')),'本輪交付不得包含 PDF')
    books={};formula_count=0
    for name in ('設計計算表.xlsx','設備明細表.xlsx'):
        values=openpyxl.load_workbook(OUT/name,data_only=True)
        formulas=openpyxl.load_workbook(OUT/name,data_only=False)
        books[name]=values
        for sheet in formulas:
            check(bool(sheet.print_area),f'{sheet.title} 缺列印範圍')
            check(bool(sheet.print_title_rows),f'{sheet.title} 缺重複欄頭')
            check(sheet.page_setup.orientation=='landscape' and sheet.page_setup.fitToWidth==1,f'{sheet.title} 列印設定不符')
            for row in sheet:
                for cell in row:
                    cached=values[sheet.title][cell.coordinate]
                    check(cached.data_type!='e',f'{sheet.title}!{cell.coordinate} 公式錯誤')
                    if cell.data_type=='f':
                        formula_count+=1
                        check(isinstance(cached.value,(int,float)),f'{sheet.title}!{cell.coordinate} 缺數值快取')
        formulas.close()
    calc=books['設計計算表.xlsx'];equip=books['設備明細表.xlsx']
    near(calc['概覽']['C20'].value,920,'總新風 L/s');near(calc['概覽']['D20'].value,87.5,'總排風 L/s')
    for cell,expected in [('M15',920),('J15',87.5),('P15',3312),('I15',315)]:near(calc['通風計算'][cell].value,expected,'新排風換算 '+cell)
    for row in range(5,8):
        area,raw,sf,sizing,density,sizing_density=[calc['系統負荷密度'].cell(row,c).value for c in range(2,8)]
        near(sf,1.15,'SF');near(sizing,raw*1.15,'單次 SF')
        near(density,raw*1000/area,'原始密度');near(sizing_density,sizing*1000/area,'含 SF 密度')
    for row in range(5,14):
        s=calc['冷負荷計算'];near(s.cell(row,9).value,s.cell(row,7).value*1000/s.cell(row,3).value,'房間密度')
        check(not isinstance(s.cell(row,6).value,(int,float)),'新風不可重複分攤為房間負荷')
    air=package['air_treatment'];fresh=calc['新風處理計算']
    for cell,expected in [('B5',air['airflow_ls']),('B12',air['raw_coil']['total_kw']),('B13',1.15),('B14',air['sizing_coil']['total_kw']),('B15',air['reheat_kw']),('B17',0)]:near(fresh[cell].value,expected,'新風 '+cell)
    near(fresh['B14'].value,fresh['B12'].value*1.15,'新風單次 SF')
    near(fresh['B12'].value-fresh['B15'].value,air['net_outdoor_to_supply_kw'],'新風焓差閉合')
    near(air['coil_outlet']['humidity_ratio_kgkg'],air['supply']['humidity_ratio_kgkg'],'再熱不改含濕量')
    calc_text,equip_text=all_text(calc),all_text(equip)
    with ZipFile(OUT/'設計說明.docx') as z:doc=ET.fromstring(z.read('word/document.xml'))
    word_text=''.join(n.text or '' for n in doc.iter(W+'t'))
    for forbidden in ('項目地點','實際工程地址','設計測試','本測試','本試點','測試假設','畫圖交接','CalculationOnly','只供','冬季','假設可信度','未計算'):
        check(forbidden not in word_text,f'正式報告仍有流程文案：{forbidden}')
    check(len(list(doc.iter(W+'tbl')))==4,'FOP 四表缺漏')
    template_path=ROOT/'templates/fop_fixed_content_v1.1.json'
    if not template_path.exists():
        # Read the builder's versioned template path; do not duplicate a second template source.
        import importlib.util
        spec=importlib.util.spec_from_file_location('selected_narrative',ROOT/'scripts/build_retail_trial_narrative.py')
        mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod);template_path=mod.TEMPLATE
    template=json.loads(template_path.read_text(encoding='utf-8'))
    check(set(template['fixed_chapters'])=={'2','4','6','7'},'固定章集合改變')
    for chapter in template['fixed_chapters'].values():
        for p in chapter:check(p['text'] in word_text,f'固定原文缺漏：{p["paragraph_index"]}')
    expected={e['id'] for e in package['equipment']}
    actual={str(row[0].value) for s in equip for row in s.iter_rows(min_row=8) if row[0].value in expected}
    check(actual==expected,'設備需求編號不完整')
    for identifier in expected:check(identifier in calc_text and identifier in word_text,f'{identifier} 跨文件對映缺漏')
    for e in package['equipment']:
        s=equip[{'AC':'空調設備','TEF':'風機','FAU':'新風處理'}[e['category']]]
        row=next(r for r in s if r[0].value==e['id'])
        near(row[3 if e['category']=='AC' else 2].value,e['quantity'],e['id']+' 數量')
        for rid in e['room_ids']:check(rid in str(row[1].value),e['id']+' 服務房間缺漏 '+rid)
        if e['category']=='AC':near(row[5].value,e['required_cooling_kw'],e['id']+' 系統需求')
        elif e['category']=='FAU':near(row[7].value,air['sizing_coil']['total_kw'],'FAU 盤管需求')
    supplement=package['source_supplement'];sources={s['id']:s for s in supplement['sources']}
    candidates={c['model']:c for c in supplement['candidates']}
    configurations=package['configuration'];check(bool(configurations),'缺機組配置')
    ids=[c['equipment_id'] for c in configurations];check(len(ids)==len(set(ids)),'設備編號重複')
    sheet=equip['機組配置']
    for c in configurations:
        check(c['model'] in candidates,'配置型號未經原始來源覆核：'+c['model'])
        candidate=candidates[c['model']]
        check(candidate['source_id'] in sources,'型號來源不存在')
        check(candidate['status'] not in ('RejectedForCurrentDuty','Flagged'),'淘汰候選被配置')
        check(c.get('review_status')!='Flagged','Flagged 型號被推薦')
        check(c['quantity']>0 and int(c['quantity'])==c['quantity'],'設備數量須正整數')
        row=next((r for r in sheet if r[0].value==c['equipment_id']),None)
        check(row is not None,'設備表缺配置 '+c['equipment_id'])
        check(c['model'] in str(row[2].value),'設備型號與配置不符')
        near(row[4].value,c['quantity'],c['equipment_id']+' 台數')
        capacity=c.get('nominal_cooling_kw')
        if capacity is not None:near(row[6].value,capacity,c['equipment_id']+' 單機標稱冷量')
        else:check(not isinstance(row[6].value,(int,float)),'未知單機冷量變成數值')
        check(c['model'] in calc_text,'計算表未列配置型號')
        for rid in c.get('room_ids',[]):check(rid in str(row[5].value),'配置服务房間不符')
    # Explicitly reject the source-name mismatch; similarity is not an approved alias mapping.
    suspect={'RAS-4HNRQ','RAS-5HNRQ','RAS-6HNRQ','RAS-8HNRQ','RAS-10HNRQ','RAS-12HNRQ'}
    check(not suspect.intersection(c['model'] for c in configurations),'資料庫型號差異未隔離')
    check(bool(supplement.get('database_quality_findings')),'原始型錄差異缺紀錄')
    snapshot=package.get('equipment_snapshot',package.get('snapshot',{}))
    check(bool(snapshot.get('queried_at')) and bool(snapshot.get('source')),'快照日期或原來源缺失')
    for unit in ('m²','m³','(人)','L/s','m³/h','kW','W/m²','Pa','ACH'):check(unit in calc_text,'計算表缺單位 '+unit)
    for unit in ('台','kW','L/s','Pa','rpm','V/Ph/Hz','dB(A)','mm','°C DB'):check(unit in equip_text,'設備表缺單位 '+unit)
    check(formula_count>0,'没有可驗收公式')
    for book in books.values():book.close()
    print(f'通過：三份成果、{formula_count} 個公式快取、{len(expected)} 個需求、{len(configurations)} 條配置；新風質能平衡、單次SF、固定文案、來源隔離及列印設定。')


if __name__=='__main__':main()
